#!/usr/bin/env python3
"""Generate synthetic binary fixtures for the resume parser benchmark.

Benchmark-only dependencies:
  python-docx==1.2.0
  reportlab==4.4.4
  Pillow==11.3.0
  pypdf==6.0.0

All people and organizations in these fixtures are fictional.
"""

from pathlib import Path
import zipfile

from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "tests" / "fixtures" / "career-evidence" / "parser-benchmark" / "generated"


def make_docx(filename: str, name: str, role: str, org: str, dates: str, bullets: list[str]) -> None:
    document = Document()
    document.add_heading(name, 0)
    document.add_paragraph(f"{role} | {org} | {dates}")
    for bullet in bullets:
        document.add_paragraph(bullet, style="List Bullet")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Customer service | Documentation | Safety")
    document.save(OUT / filename)


def draw_text_pdf(filename: str, lines: list[tuple[str, int]]) -> None:
    output = canvas.Canvas(str(OUT / filename), pagesize=LETTER, invariant=1)
    y = 750
    for text, size in lines:
        output.setFont("Helvetica-Bold" if size >= 11 else "Helvetica", size)
        output.drawString(54, y, text)
        y -= 24 if size >= 11 else 18
    output.save()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    make_docx(
        "hvac-technician.docx",
        "Morgan Reed",
        "HVAC Service Technician",
        "Harbor Mechanical",
        "2018-2026",
        [
            "EPA 608 Universal certification",
            "Serviced rooftop units, heat pumps, split systems, and gas furnaces",
            "Performed electrical troubleshooting and preventive maintenance",
            "Participated in rotating on-call coverage",
        ],
    )

    make_docx(
        "registered-nurse.docx",
        "Casey Lin",
        "Registered Nurse",
        "Cedar Point Medical Center",
        "2020-present",
        [
            "Active RN license",
            "BLS certification",
            "Medical-surgical unit experience",
            "Precepted newly hired nurses",
        ],
    )

    draw_text_pdf(
        "operations-history.pdf",
        [
            ("Samira Cole", 16),
            ("Administrative Coordinator | Bluewater Association | 2022-present", 11),
            ("• Managed executive calendars and meeting logistics", 10),
            ("• Processed vendor invoices and purchase requests", 10),
            ("• Created monthly operational reports in Excel", 10),
            ("Office Assistant | Harbor Services | 2019-2022", 11),
            ("• Responded to customer inquiries and maintained member records", 10),
            ("• Coordinated meeting rooms and supply orders", 10),
        ],
    )

    draw_text_pdf(
        "career-changer.pdf",
        [
            ("Elena Park", 16),
            ("High School Teacher | East Ridge School | 2016-2026", 11),
            ("• Designed curriculum and facilitated professional-development sessions", 10),
            ("• Coordinated a six-person curriculum committee", 10),
            ("• Built asynchronous learning materials in the district LMS", 10),
        ],
    )

    (OUT / "federal-applicant.txt").write_text(
        "Cameron West\n"
        "Program Specialist | Pine County Government | 2021-present | 40 hours/week\n"
        "- Administered public-service program intake and eligibility records\n"
        "- Prepared quarterly compliance reports\n"
        "- Trained six staff members on revised procedures\n",
        encoding="utf-8",
    )

    image = Image.new("RGB", (1650, 2200), "white")
    drawing = ImageDraw.Draw(image)
    drawing.multiline_text(
        (100, 120),
        "Riley Brooks\nWarehouse Supervisor\nForklift operator certification\n"
        "Supervised 18 associates across receiving and fulfillment",
        fill="black",
        spacing=18,
    )
    image.save(OUT / "scanned-image-only.pdf", "PDF", resolution=150.0)

    reader = PdfReader(str(OUT / "career-changer.pdf"))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("benchmark-password")
    with (OUT / "encrypted-resume.pdf").open("wb") as handle:
        writer.write(handle)

    with zipfile.ZipFile(OUT / "malformed.docx", "w") as archive:
        archive.writestr("broken.txt", "this is not a Word document package")

    (OUT / "unsupported.bin").write_bytes(b"\x00\x01JOBRANGER\xff\xfeunsupported-format")

    print(f"Generated parser fixtures in {OUT}")


if __name__ == "__main__":
    main()
